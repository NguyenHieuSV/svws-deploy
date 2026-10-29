"""
Mẫu BOQ 12 cột (29/10/2026): thêm Nhà cung cấp · Mã SP; nạp file (mẫu BOQ hoặc AI) gắn NCC trong hồ sơ và ghi ngược danh mục
Sản phẩm NCC có kiểm soát; mẫu cũ 10 cột vẫn đọc được. SQLite trong bộ nhớ.
"""
import io
from datetime import date
from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker

from app import boq_mau as bm
from app import models as m
from app.routers import ncc

NV = SimpleNamespace(id=1)          # người nạp (audit đã chặn trong fixture)
NCC_TEN = "Công ty TNHH Tự Động Hóa Thành Công"


def test_tao_mau_12_cot_cong_thuc():
    ws = load_workbook(io.BytesIO(bm.tao_mau("SVWS", "Thử", "DT-TEST", None))).active
    assert [ws.cell(row=4, column=j).value for j in range(1, 13)] == bm.COT
    assert ws.cell(row=4, column=7).value == "Nhà cung cấp" and ws.cell(row=4, column=8).value == "Mã SP"
    assert ws["A6"].value == "A01" and ws["K6"].value == "=E6*J6"                      # thành tiền = KL × Đơn giá (E × J)
    tong = [c for row in ws.iter_rows(min_col=2, max_col=2) for c in row if c.value == "CỘNG TRƯỚC THUẾ"][0]
    assert str(ws.cell(row=tong.row, column=11).value).startswith("=SUM(K")           # tổng theo cột K


def test_xuat_roi_nap_lai_giu_ncc_ma_sp():
    items = [{"ten": "Bơm định lượng", "quy_cach": "0–10 l/h", "don_vi": "cái", "so_luong": 2, "nha_san_xuat": "Pulsafeeder",
              "ncc_ten": "Công ty ABC", "ma_sp": "PF-10", "don_gia": 5_000_000,
              "ghi_chu": "BOQ A01 · nhóm A THIẾT BỊ · Nguồn: Đề xuất"}]
    nhom, ds = bm.doc_boq(bm.tao_mau("SVWS", "Thử", "DT-1", items), "x.xlsx")
    assert [n["ma"] for n in nhom] == ["A"] and len(ds) == 1
    r = ds[0]
    assert r["ncc_ten"] == "Công ty ABC" and r["ma_sp"] == "PF-10" and r["ma"] == "A01"
    assert r["so_luong"] == 2 and r["don_gia"] == 5_000_000 and r["nha_san_xuat"] == "Pulsafeeder"


def test_mau_cu_10_cot_van_doc():
    csv = ("Mã,Hạng mục,Quy cách,ĐVT,KL,Nhà sản xuất,Nguồn,Đơn giá (VND),Thành tiền (VND),Căn cứ tính\n"
           "A,THIẾT BỊ,,,,,,,,\n"
           "A01,Cột lọc,Ø500,cái,1,SVWS,Đề xuất,1000000,1000000,RFQ\n").encode("utf-8")
    nhom, ds = bm.doc_boq(csv, "boq.csv")
    assert len(nhom) == 1 and len(ds) == 1
    assert ds[0]["ncc_ten"] == "" and ds[0]["ma_sp"] == "" and ds[0]["don_gia"] == 1_000_000


# ---------- nạp vào dự toán: gắn NCC + ghi ngược danh mục ----------
@pytest.fixture()
def db(monkeypatch):
    monkeypatch.setattr(ncc, "ghi_audit", lambda *a, **k: None)
    eng = create_engine("sqlite://")
    for t in (m.NhaCungCap, m.SanPhamNcc, m.DuToanBan, m.DuToanBanMuc, m.HangHoa, m.TonKho):
        t.__table__.create(eng)
    s = sessionmaker(eng)()
    n = {}

    @event.listens_for(s, "before_flush")
    def _ids(sess, _c, _i):
        for o in sess.new:
            if getattr(o, "id", None) is None:
                n[type(o)] = n.get(type(o), 0) + 1
                o.id = n[type(o)]

    s.add(m.NhaCungCap(ten=NCC_TEN)); s.flush()
    s.add(m.SanPhamNcc(nha_cung_cap_id=1, ten="Tủ điện EDI", ma_sp="TD-EDI", don_vi="bộ", don_gia=Decimal(80_000_000)))
    s.add(m.DuToanBan(ma="DT-TEST", ngay=date(2026, 10, 29)))
    s.commit()
    return s


def test_nap_mau_boq_gan_ncc_va_danh_muc(db):
    dt = db.get(m.DuToanBan, 1)
    items = [
        {"ma": "A01", "ten": "Tủ điện EDI", "so_luong": 1, "don_gia": 80_000_000, "ncc_ten": NCC_TEN.lower(), "ma_sp": "TD-EDI"},  # trùng mã → nối
        {"ma": "A02", "ten": "Tủ điều khiển RO", "so_luong": 1, "don_gia": 45_000_000, "ncc_ten": NCC_TEN},                    # mới → tạo SP
        {"ma": "A03", "ten": "Bơm cấp", "so_luong": 2, "don_gia": 3_000_000, "ncc_ten": "NCC Lạ"},                               # không có hồ sơ
        {"ma": "A04", "ten": "Van tay", "so_luong": 5, "don_gia": 100_000},                                                      # không ghi NCC
    ]
    kq = ncc._dtb_them_muc_boq(db, NV, dt, items, "boq.xlsx")
    db.commit()
    assert kq["them"] == 4 and kq["ncc_gan"] == 2 and kq["dm_noi"] == 1 and kq["dm_tao"] == 1 and kq["ncc_la"] == ["NCC Lạ"]
    rows = {r.ten: r for r in db.query(m.DuToanBanMuc).all()}
    assert rows["Tủ điện EDI"].san_pham_ncc_id == 1 and rows["Tủ điện EDI"].nha_cung_cap_id == 1
    sp_moi = db.query(m.SanPhamNcc).filter_by(ten="Tủ điều khiển RO").one()
    assert rows["Tủ điều khiển RO"].san_pham_ncc_id == sp_moi.id and float(sp_moi.don_gia) == 45_000_000
    assert "Từ dự toán DT-TEST" in sp_moi.ghi_chu
    assert rows["Bơm cấp"].nha_cung_cap_id is None and rows["Bơm cấp"].ncc_ten == "NCC Lạ" and rows["Bơm cấp"].san_pham_ncc_id is None
    assert rows["Van tay"].ncc_ten is None and rows["Van tay"].nha_cung_cap_id is None
    assert float(db.get(m.SanPhamNcc, 1).don_gia) == 80_000_000                       # giá danh mục không bị ghi đè


def test_ai_items_gan_ncc_va_tach_cot(db):
    items = [{"ten": "Tủ điện EDI", "ma_sp": "TD-EDI", "mo_ta": "Tủ inox", "nha_san_xuat": "Schneider", "spec": "380V; IP54",
              "nha_cung_cap": NCC_TEN, "don_vi": "bộ", "so_luong": 1, "don_gia": 80_000_000},
             {"ten": "Ống PVC", "nha_cung_cap": None, "don_vi": "m", "so_luong": 10, "don_gia": 20_000}]
    kq = ncc._dtb_them_muc_tu_items(db, 1, items, "bg.pdf", nd=NV, ma_dt="DT-TEST")
    db.commit()
    assert kq["them"] == 2 and kq["ncc_gan"] == 1 and kq["dm_noi"] == 1 and kq["dm_tao"] == 0 and kq["ncc_la"] == []
    r = db.query(m.DuToanBanMuc).filter_by(ten="Tủ điện EDI").one()
    assert r.san_pham_ncc_id == 1 and r.nha_cung_cap_id == 1 and r.ncc_ten == NCC_TEN
    assert r.quy_cach == "Tủ inox" and r.spec == "380V; IP54" and r.nha_san_xuat == "Schneider" and r.ma_sp == "TD-EDI"
    r2 = db.query(m.DuToanBanMuc).filter_by(ten="Ống PVC").one()
    assert r2.nha_cung_cap_id is None and r2.ncc_ten is None and r2.san_pham_ncc_id is None
    # không có người nạp (nd=None) → chỉ gắn NCC, không ghi danh mục
    kq2 = ncc._dtb_them_muc_tu_items(db, 1, [{"ten": "Bơm RO", "nha_cung_cap": NCC_TEN, "so_luong": 1, "don_gia": 9_000_000}], "x.pdf")
    assert kq2["ncc_gan"] == 1 and kq2["dm_tao"] == 0 and db.query(m.SanPhamNcc).count() == 1
