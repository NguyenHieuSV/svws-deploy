"""
Test TRAINING (mục Nội bộ — app/routers/dao_tao.py): tải bài, xem tệp, ghi điểm, quyền, bảng điểm, Excel.
KHÔNG cần PostgreSQL: dùng SQLite trong bộ nhớ (JSONB → JSON), thay phân quyền / audit bằng bản giả.
"""
import asyncio
import io
import os
import tempfile

import pytest
from fastapi import HTTPException, UploadFile
from sqlalchemy import create_engine, event
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import sessionmaker


@compiles(JSONB, "sqlite")
def _jsonb_sqlite(t, c, **k):
    return "JSON"


os.environ.setdefault("STORAGE_DIR", tempfile.mkdtemp(prefix="trn_test_"))
from app import models as m                      # noqa: E402
from app.routers import dao_tao as d             # noqa: E402

HTML = "<html><head><title>x</title></head><body>bài thi</body></html>".encode("utf-8")


@pytest.fixture()
def ctx(monkeypatch):
    muc = {"v": "XEM"}
    monkeypatch.setattr(d, "ghi_audit", lambda *a, **k: None)
    monkeypatch.setattr(d, "muc_cua", lambda db, vt, mod: muc["v"])
    eng = create_engine("sqlite://")
    for t in (m.VaiTro, m.NguoiDung, m.NhanVien, m.TrnBai, m.TrnKetQua):
        t.__table__.create(eng)
    db = sessionmaker(eng)()
    n = {}

    @event.listens_for(db, "before_flush")
    def _ids(sess, _ctx, _inst):
        for o in sess.new:
            if getattr(o, "id", None) is None:
                n[type(o)] = n.get(type(o), 0) + 1
                o.id = n[type(o)]

    db.add_all([m.VaiTro(id=1, ma="CEO", ten="CEO"), m.VaiTro(id=2, ma="NV_KT", ten="NV")])
    db.flush()
    ql = m.NguoiDung(vai_tro_id=1, email="ceo@x.vn", mat_khau_hash="x", trang_thai="HOAT_DONG")
    nv = m.NguoiDung(vai_tro_id=2, email="nv@x.vn", mat_khau_hash="x", trang_thai="HOAT_DONG")
    db.add_all([ql, nv])
    db.flush()
    db.add_all([m.NhanVien(nguoi_dung_id=nv.id, ma="NV01", ho_ten="Nguyễn Văn A", chuc_danh="Vận hành", trang_thai="DANG_LAM"),
                m.NhanVien(ma="NV02", ho_ten="Trần Thị B", trang_thai="DANG_LAM")])
    db.commit()
    return db, ql, nv, muc


def _up(ten, nd):
    return UploadFile(filename=ten, file=io.BytesIO(nd))


def _them(db, nd, **kw):
    a = dict(tieu_de="Đào tạo EC", loai="HUAN_LUYEN_THI", mo_ta=None, diem_dat=80.0, hien=True, file=_up("bai.html", HTML), db=db, nd=nd)
    a.update(kw)
    return asyncio.run(d.them_bai(**a))


def test_tai_bai_va_chan_dinh_dang(ctx):
    db, ql, nv, muc = ctx
    muc["v"] = "THAO_TAC"
    b = _them(db, ql)
    assert b["loai_tep"] == "html" and b["phien_ban"] == 1 and b["co_thi"] is True
    for kw in ({"file": _up("a.exe", b"MZ")}, {"loai": "SAI"}, {"tieu_de": "  "}, {"diem_dat": 120.0}):
        with pytest.raises(HTTPException) as e:
            _them(db, ql, **kw)
        assert e.value.status_code == 400


def test_an_bai_va_xem_tep(ctx):
    db, ql, nv, muc = ctx
    muc["v"] = "THAO_TAC"
    b1, b2 = _them(db, ql), _them(db, ql, tieu_de="Ẩn", hien=False)
    muc["v"] = "XEM"
    assert [x["id"] for x in d.ds_bai(db, nv)] == [b1["id"]]          # nhân viên không thấy bài ẩn
    with pytest.raises(HTTPException) as e:
        d.xem_tep_bai(b2["id"], db, nv)
    assert e.value.status_code == 404
    r = d.xem_tep_bai(b1["id"], db, nv)
    assert r.body == HTML and "sandbox" in r.headers["content-security-policy"]
    muc["v"] = "THAO_TAC"
    assert {x["id"] for x in d.ds_bai(db, ql)} == {b1["id"], b2["id"]}


def test_ghi_diem_va_phan_quyen_ket_qua(ctx):
    db, ql, nv, muc = ctx
    muc["v"] = "THAO_TAC"
    b = _them(db, ql)
    muc["v"] = "XEM"
    k1 = d.ghi_ket_qua(b["id"], d.KetQuaVao(diem=86.7, so_cau=30, so_dung=26, ten_khai="A", chi_tiet={"det": []}), db, nv)
    assert k1["lan_thi"] == 1 and k1["dat"] is True and k1["ten"] == "Nguyễn Văn A" and k1["nhan_vien_id"]
    assert d.ghi_ket_qua(b["id"], d.KetQuaVao(diem=86.7, so_cau=30, so_dung=26), db, nv)["id"] == k1["id"]   # gửi lặp ≤ 20 giây
    k2 = d.ghi_ket_qua(b["id"], d.KetQuaVao(diem=60, so_cau=30, so_dung=18), db, nv)
    assert k2["lan_thi"] == 2 and k2["dat"] is False                                         # tự tính theo ngưỡng 80 của bài
    with pytest.raises(HTTPException) as e:
        d.ghi_ket_qua(b["id"], d.KetQuaVao(diem=50, so_cau=3, so_dung=5), db, nv)
    assert e.value.status_code == 400
    muc["v"] = "THAO_TAC"
    k3 = d.ghi_ket_qua(b["id"], d.KetQuaVao(diem=95, so_cau=30, so_dung=29), db, ql)
    muc["v"] = "XEM"
    assert {x["id"] for x in d.ds_ket_qua(None, None, None, False, 500, db, nv)} == {k1["id"], k2["id"]}   # chỉ thấy của mình
    with pytest.raises(HTTPException) as e:
        d.chi_tiet_ket_qua(k3["id"], db, nv)
    assert e.value.status_code == 404
    muc["v"] = "THAO_TAC"
    assert len(d.ds_ket_qua(None, None, None, False, 500, db, ql)) == 3
    cua_toi = [x for x in d.ds_bai(db, nv)][0]["cua_toi"]
    assert cua_toi["so_lan"] == 2 and cua_toi["tot_nhat"] == 86.7 and cua_toi["dat"] is True


def test_bang_diem_excel_thay_tep_xoa(ctx):
    db, ql, nv, muc = ctx
    muc["v"] = "THAO_TAC"
    b = _them(db, ql)
    muc["v"] = "XEM"
    d.ghi_ket_qua(b["id"], d.KetQuaVao(diem=90, so_cau=30, so_dung=27), db, nv)
    bd = d._bang_diem(db)
    cac = {h["ten"]: h for h in bd["nhan_vien"]}
    assert cac["Nguyễn Văn A"]["so_bai_dat"] == 1 and cac["Trần Thị B"]["so_bai_thi"] == 0   # chưa thi vẫn có dòng
    from openpyxl import load_workbook
    rows = db.query(m.TrnKetQua, m.TrnBai.tieu_de).join(m.TrnBai, m.TrnBai.id == m.TrnKetQua.bai_id).all()
    wb = load_workbook(io.BytesIO(d._excel_bytes(bd, [d._kq_dict(r, t) for r, t in rows])))
    assert wb.sheetnames == ["Bảng điểm", "Lượt thi"]
    muc["v"] = "THAO_TAC"
    t = asyncio.run(d.thay_tep_bai(b["id"], _up("RevF.html", HTML + b"<!--F-->"), db, ql))
    assert t["phien_ban"] == 2
    kq = d.ghi_ket_qua(b["id"], d.KetQuaVao(diem=70, so_cau=30, so_dung=21), db, ql)
    assert kq["phien_ban"] == 2                                                              # điểm ghi rõ thi ở Rev nào
    muc["v"] = "QUAN_TRI"
    with pytest.raises(HTTPException) as e:
        d.xoa_bai(b["id"], db, ql)
    assert e.value.status_code == 409                                                        # đã có điểm → chỉ được ẩn
    b3 = _them(db, ql, tieu_de="Chưa ai thi")
    assert d.xoa_bai(b3["id"], db, ql) == {"ok": True}
