"""TRAINING (mục Nội bộ) — kho bài huấn luyện / bài thi + ĐIỂM THI của từng nhân viên (mig 149).

• Bài = một tệp (HTML tương tác tự chấm điểm, PDF, Office, ảnh, video…) lưu ở kho lưu trữ chung (app/luu_tru.py), có phiên bản (Rev).
• Bài HTML chạy trong khung cách ly (sandbox) ở giao diện; khi người dùng nộp bài, khung báo điểm về app → POST /dao-tao/bai/{id}/ket-qua.
  Điểm LUÔN gắn với người đang đăng nhập (không tin tên gõ trong bài; tên gõ chỉ lưu thêm để đối chiếu).
• Quyền (module 'training'): XEM = học, thi, xem điểm của mình · THAO_TAC = tải / sửa / ẩn bài, xem bảng điểm mọi người, xuất Excel ·
  QUAN_TRI = xóa bài (chưa có kết quả), xóa kết quả thi."""
import io
import json
from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Form
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session
from ..database import get_db
from ..rbac import yeu_cau, muc_cua, THU_BAC
from ..audit import ghi_audit
from ..models import NguoiDung, NhanVien, TrnBai, TrnKetQua
from ..luu_tru import luu, xoa, doc, ton_tai, phan_hoi_tai
from ..nhac_viec_service import gio_hien_tai

router = APIRouter(prefix="/dao-tao", tags=["dao_tao"])
MODULE = "training"

LOAI_BAI = {"HUAN_LUYEN": "Huấn luyện", "THI": "Thi", "HUAN_LUYEN_THI": "Huấn luyện + thi"}
LOAI_CO_THI = ("THI", "HUAN_LUYEN_THI")
GIOI_HAN_BYTE = 60 * 1024 * 1024
# đuôi tệp cho phép → (loại hiển thị, content-type)
DUOI_TEP = {
    "html": ("html", "text/html; charset=utf-8"), "htm": ("html", "text/html; charset=utf-8"),
    "pdf": ("pdf", "application/pdf"),
    "png": ("hinh", "image/png"), "jpg": ("hinh", "image/jpeg"), "jpeg": ("hinh", "image/jpeg"), "gif": ("hinh", "image/gif"),
    "mp4": ("video", "video/mp4"), "webm": ("video", "video/webm"), "mp3": ("video", "audio/mpeg"),
    "ppt": ("khac", "application/vnd.ms-powerpoint"),
    "pptx": ("khac", "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
    "doc": ("khac", "application/msword"),
    "docx": ("khac", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "xls": ("khac", "application/vnd.ms-excel"),
    "xlsx": ("khac", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
}


def _la_quan_ly(db: Session, nd: NguoiDung) -> bool:
    return THU_BAC.get(muc_cua(db, nd.vai_tro_id, MODULE), 0) >= THU_BAC["THAO_TAC"]


def _nv_cua(db: Session, nd: NguoiDung):
    return (db.query(NhanVien).filter(NhanVien.nguoi_dung_id == nd.id).order_by(NhanVien.id).first())


def _ten_nguoi(db: Session, nd: NguoiDung) -> str:
    nv = _nv_cua(db, nd)
    return (nv.ho_ten if nv is not None and nv.ho_ten else nd.email) or f"ND#{nd.id}"


def _duoi(ten_file: str) -> str:
    return (ten_file or "").rsplit(".", 1)[-1].lower() if "." in (ten_file or "") else ""


def _f(v):
    return float(v) if v is not None else None


def _ten_tao(db: Session, nd_id):
    if not nd_id:
        return None
    nd = db.get(NguoiDung, nd_id)
    if nd is None:
        return None
    return _ten_nguoi(db, nd)


def _bai_dict(b: TrnBai, ten_tao=None, thong_ke=None, cua_toi=None) -> dict:
    return {"id": b.id, "tieu_de": b.tieu_de, "loai": b.loai, "loai_ten": LOAI_BAI.get(b.loai, b.loai), "mo_ta": b.mo_ta,
            "ten_file": b.ten_file, "loai_tep": b.loai_tep, "kich_thuoc": int(b.kich_thuoc or 0), "phien_ban": b.phien_ban or 1,
            "hien": bool(b.hien), "diem_dat": _f(b.diem_dat), "nguoi_tao": ten_tao,
            "tao_luc": str(b.tao_luc)[:16] if b.tao_luc else None,
            "cap_nhat_luc": str(b.cap_nhat_luc)[:16] if b.cap_nhat_luc else None,
            "co_thi": b.loai in LOAI_CO_THI,
            "thong_ke": thong_ke, "cua_toi": cua_toi}


def _kq_dict(r: TrnKetQua, ten_bai=None, chi_tiet=False) -> dict:
    o = {"id": r.id, "bai_id": r.bai_id, "bai": ten_bai, "nguoi_dung_id": r.nguoi_dung_id, "nhan_vien_id": r.nhan_vien_id,
         "ten": r.ten_nguoi_thi, "ten_khai": r.ten_khai, "ma_nv_khai": r.ma_nv_khai, "don_vi_khai": r.don_vi_khai,
         "chuc_danh_khai": r.chuc_danh_khai, "diem": _f(r.diem), "so_cau": r.so_cau, "so_dung": r.so_dung,
         "diem_dat": _f(r.diem_dat), "dat": r.dat, "lan_thi": r.lan_thi, "phien_ban": r.phien_ban,
         "thoi_gian_giay": r.thoi_gian_giay, "nop_luc": str(r.nop_luc)[:16] if r.nop_luc else None}
    if chi_tiet:
        o["chi_tiet"] = r.chi_tiet
    return o


# ================================================================== BÀI HUẤN LUYỆN / THI
@router.get("/toi")
def thong_tin_toi(db: Session = Depends(get_db), nd: NguoiDung = Depends(yeu_cau(MODULE, "XEM"))):
    """Thông tin người đang đăng nhập — điền sẵn vào bài thi."""
    nv = _nv_cua(db, nd)
    return {"ten": _ten_nguoi(db, nd), "ma_nv": nv.ma if nv is not None else None,
            "chuc_danh": nv.chuc_danh if nv is not None else None, "email": nd.email,
            "quan_ly": _la_quan_ly(db, nd)}


@router.get("/bai")
def ds_bai(db: Session = Depends(get_db), nd: NguoiDung = Depends(yeu_cau(MODULE, "XEM"))):
    """Danh sách bài (người quản lý thấy cả bài đang ẩn) + thống kê lượt thi + điểm của chính người xem."""
    ql = _la_quan_ly(db, nd)
    q = db.query(TrnBai)
    if not ql:
        q = q.filter(TrnBai.hien.is_(True))
    bais = q.order_by(TrnBai.id.desc()).all()
    ids = [b.id for b in bais]
    tk, ct = {}, {}
    if ids:
        for bid, nvid, ndid, diem, dat, nop in (db.query(TrnKetQua.bai_id, TrnKetQua.nhan_vien_id, TrnKetQua.nguoi_dung_id,
                                                         TrnKetQua.diem, TrnKetQua.dat, TrnKetQua.nop_luc)
                                                .filter(TrnKetQua.bai_id.in_(ids)).all()):
            t = tk.setdefault(bid, {"so_luot": 0, "_nguoi": set(), "so_dat": 0})
            t["so_luot"] += 1
            t["_nguoi"].add(("nv", nvid) if nvid else ("nd", ndid))
            if dat:
                t["so_dat"] += 1
            if ndid == nd.id:
                c = ct.setdefault(bid, {"so_lan": 0, "tot_nhat": None, "gan_nhat": None, "dat": False, "ngay": None, "_nop": None})
                c["so_lan"] += 1
                d = float(diem)
                c["tot_nhat"] = d if c["tot_nhat"] is None else max(c["tot_nhat"], d)
                if dat:
                    c["dat"] = True
                if c["_nop"] is None or (nop and nop >= c["_nop"]):
                    c["_nop"], c["gan_nhat"], c["ngay"] = nop, d, (str(nop)[:10] if nop else None)
    out = []
    for b in bais:
        t = tk.get(b.id)
        thong_ke = ({"so_luot": t["so_luot"], "so_nguoi": len(t["_nguoi"]), "so_dat": t["so_dat"]} if t
                    else {"so_luot": 0, "so_nguoi": 0, "so_dat": 0})
        c = ct.get(b.id)
        if c:
            c.pop("_nop", None)
        out.append(_bai_dict(b, _ten_tao(db, b.nguoi_tao), thong_ke if ql else None, c))
    return out


async def _doc_tep_tai_len(file: UploadFile):
    if file is None or not file.filename:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Chưa chọn tệp bài huấn luyện")
    duoi = _duoi(file.filename)
    if duoi not in DUOI_TEP:
        raise HTTPException(status.HTTP_400_BAD_REQUEST,
                            "Định dạng chưa hỗ trợ. Dùng: " + ", ".join(sorted(DUOI_TEP)))
    data = await file.read()
    if not data:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Tệp rỗng")
    if len(data) > GIOI_HAN_BYTE:
        raise HTTPException(status.HTTP_400_BAD_REQUEST,
                            f"Tệp quá lớn ({len(data) / 1048576:.1f} MB) — tối đa {GIOI_HAN_BYTE // 1048576} MB")
    loai_tep, ctype = DUOI_TEP[duoi]
    return data, loai_tep, ctype


@router.post("/bai", status_code=201)
async def them_bai(tieu_de: str = Form(...), loai: str = Form("HUAN_LUYEN"), mo_ta: str = Form(None),
                   diem_dat: float = Form(None), hien: bool = Form(True), file: UploadFile = File(...),
                   db: Session = Depends(get_db), nd: NguoiDung = Depends(yeu_cau(MODULE, "THAO_TAC"))):
    """Tải LÊN một bài huấn luyện / bài thi."""
    if loai not in LOAI_BAI:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Loại bài không hợp lệ")
    tieu_de = (tieu_de or "").strip()
    if not tieu_de:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Nhập tiêu đề bài")
    if diem_dat is not None and not (0 <= diem_dat <= 100):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Điểm đạt phải từ 0 đến 100 (%)")
    data, loai_tep, ctype = await _doc_tep_tai_len(file)
    ref = luu(data, "dao_tao", "bai", file.filename, ctype)
    b = TrnBai(tieu_de=tieu_de[:250], loai=loai, mo_ta=(mo_ta or "").strip() or None, ten_file=file.filename[:250],
               loai_tep=loai_tep, noi_dung_loai=ctype[:120], duong_dan=ref, kich_thuoc=len(data), phien_ban=1,
               hien=bool(hien), diem_dat=diem_dat, nguoi_tao=nd.id, tao_luc=gio_hien_tai())
    db.add(b)
    db.flush()
    ghi_audit(db, nd.id, "TAO", "trn_bai", b.id, moi={"tieu_de": b.tieu_de, "loai": loai, "file": b.ten_file, "kb": round(len(data) / 1024)})
    db.commit()
    return _bai_dict(b, _ten_nguoi(db, nd))


class BaiSua(BaseModel):
    tieu_de: str | None = None
    loai: str | None = None
    mo_ta: str | None = None
    diem_dat: float | None = Field(None, ge=0, le=100)
    xoa_diem_dat: bool = False
    hien: bool | None = None


@router.put("/bai/{bai_id}")
def sua_bai(bai_id: int, data: BaiSua, db: Session = Depends(get_db), nd: NguoiDung = Depends(yeu_cau(MODULE, "THAO_TAC"))):
    b = db.get(TrnBai, bai_id)
    if b is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Không tìm thấy bài")
    cu = {"tieu_de": b.tieu_de, "loai": b.loai, "hien": b.hien, "diem_dat": _f(b.diem_dat)}
    if data.tieu_de is not None:
        if not data.tieu_de.strip():
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Tiêu đề không được trống")
        b.tieu_de = data.tieu_de.strip()[:250]
    if data.loai is not None:
        if data.loai not in LOAI_BAI:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Loại bài không hợp lệ")
        b.loai = data.loai
    if data.mo_ta is not None:
        b.mo_ta = data.mo_ta.strip() or None
    if data.xoa_diem_dat:
        b.diem_dat = None
    elif data.diem_dat is not None:
        b.diem_dat = data.diem_dat
    if data.hien is not None:
        b.hien = bool(data.hien)
    b.cap_nhat_luc = gio_hien_tai()
    ghi_audit(db, nd.id, "SUA", "trn_bai", b.id, cu=cu,
              moi={"tieu_de": b.tieu_de, "loai": b.loai, "hien": b.hien, "diem_dat": _f(b.diem_dat)})
    db.commit()
    return _bai_dict(b, _ten_tao(db, b.nguoi_tao))


@router.post("/bai/{bai_id}/thay-tep")
async def thay_tep_bai(bai_id: int, file: UploadFile = File(...), db: Session = Depends(get_db),
                       nd: NguoiDung = Depends(yeu_cau(MODULE, "THAO_TAC"))):
    """Thay tệp bằng phiên bản mới (Rev mới) — kết quả thi cũ giữ nguyên, ghi rõ thi ở phiên bản nào."""
    b = db.get(TrnBai, bai_id)
    if b is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Không tìm thấy bài")
    data, loai_tep, ctype = await _doc_tep_tai_len(file)
    cu_ref, cu_ten = b.duong_dan, b.ten_file
    b.duong_dan = luu(data, "dao_tao", "bai", file.filename, ctype)
    b.ten_file, b.loai_tep, b.noi_dung_loai, b.kich_thuoc = file.filename[:250], loai_tep, ctype[:120], len(data)
    b.phien_ban = int(b.phien_ban or 1) + 1
    b.cap_nhat_luc = gio_hien_tai()
    ghi_audit(db, nd.id, "SUA", "trn_bai", b.id, cu={"file": cu_ten, "phien_ban": b.phien_ban - 1},
              moi={"file": b.ten_file, "phien_ban": b.phien_ban, "kb": round(len(data) / 1024)})
    db.commit()
    if cu_ref:
        try:
            xoa(cu_ref)
        except Exception:
            pass
    return _bai_dict(b, _ten_tao(db, b.nguoi_tao))


def _bai_xem_duoc(db: Session, nd: NguoiDung, bai_id: int) -> TrnBai:
    b = db.get(TrnBai, bai_id)
    if b is None or (not b.hien and not _la_quan_ly(db, nd)):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Không tìm thấy bài")
    return b


@router.get("/bai/{bai_id}/tep")
def xem_tep_bai(bai_id: int, db: Session = Depends(get_db), nd: NguoiDung = Depends(yeu_cau(MODULE, "XEM"))):
    """Nội dung tệp để XEM tại chỗ (giao diện tải bằng fetch có token rồi nhúng vào khung cách ly).
    Gắn CSP sandbox: dù ai mở thẳng đường link này, tệp HTML cũng không đọc được dữ liệu của app."""
    b = _bai_xem_duoc(db, nd, bai_id)
    if not b.duong_dan or not ton_tai(b.duong_dan):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Tệp bài không còn trên kho lưu trữ — hãy tải lên lại")
    try:
        noi_dung = doc(b.duong_dan)
    except FileNotFoundError:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Tệp bài không còn trên kho lưu trữ — hãy tải lên lại")
    return Response(content=noi_dung, media_type=b.noi_dung_loai or "application/octet-stream",
                    headers={"Content-Security-Policy": "sandbox allow-scripts allow-modals allow-forms",
                             "X-Content-Type-Options": "nosniff", "Cache-Control": "no-store"})


@router.get("/bai/{bai_id}/tai-ve")
def tai_ve_bai(bai_id: int, db: Session = Depends(get_db), nd: NguoiDung = Depends(yeu_cau(MODULE, "XEM"))):
    b = _bai_xem_duoc(db, nd, bai_id)
    if not b.duong_dan or not ton_tai(b.duong_dan):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Tệp bài không còn trên kho lưu trữ — hãy tải lên lại")
    return phan_hoi_tai(b.duong_dan, b.ten_file, b.noi_dung_loai)


@router.delete("/bai/{bai_id}")
def xoa_bai(bai_id: int, db: Session = Depends(get_db), nd: NguoiDung = Depends(yeu_cau(MODULE, "QUAN_TRI"))):
    b = db.get(TrnBai, bai_id)
    if b is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Không tìm thấy bài")
    so_kq = db.query(TrnKetQua).filter(TrnKetQua.bai_id == bai_id).count()
    if so_kq:
        raise HTTPException(status.HTTP_409_CONFLICT,
                            f"Bài đã có {so_kq} lượt thi — không xóa để giữ điểm của nhân viên. Hãy chuyển sang «Ẩn» "
                            "(nhân viên không thấy nữa, điểm vẫn còn trong bảng điểm).")
    ref, ten = b.duong_dan, b.tieu_de
    ghi_audit(db, nd.id, "XOA", "trn_bai", b.id, cu={"tieu_de": ten, "file": b.ten_file})
    db.delete(b)
    db.commit()
    if ref:
        try:
            xoa(ref)
        except Exception:
            pass
    return {"ok": True}


# ================================================================== KẾT QUẢ THI
class KetQuaVao(BaseModel):
    diem: float = Field(..., ge=0, le=100)                 # điểm quy về %
    so_cau: int | None = Field(None, ge=0, le=2000)
    so_dung: int | None = Field(None, ge=0, le=2000)
    dat: bool | None = None
    diem_dat: float | None = Field(None, ge=0, le=100)
    thoi_gian_giay: int | None = Field(None, ge=0, le=172800)
    ten_khai: str | None = Field(None, max_length=150)
    ma_nv_khai: str | None = Field(None, max_length=40)
    don_vi_khai: str | None = Field(None, max_length=120)
    chuc_danh_khai: str | None = Field(None, max_length=120)
    chi_tiet: dict | list | None = None


@router.post("/bai/{bai_id}/ket-qua", status_code=201)
def ghi_ket_qua(bai_id: int, data: KetQuaVao, db: Session = Depends(get_db), nd: NguoiDung = Depends(yeu_cau(MODULE, "XEM"))):
    """Ghi MỘT LƯỢT THI cho người đang đăng nhập (bài HTML tự báo điểm về khi nộp bài)."""
    b = _bai_xem_duoc(db, nd, bai_id)
    if data.so_cau is not None and data.so_dung is not None and data.so_dung > data.so_cau:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Số câu đúng không thể lớn hơn tổng số câu")
    if data.chi_tiet is not None and len(json.dumps(data.chi_tiet, ensure_ascii=False)) > 120_000:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Chi tiết bài thi quá lớn")
    now = gio_hien_tai()
    truoc = (db.query(TrnKetQua).filter(TrnKetQua.bai_id == bai_id, TrnKetQua.nguoi_dung_id == nd.id)
             .order_by(TrnKetQua.id.desc()).all())
    if truoc and truoc[0].nop_luc and (now - truoc[0].nop_luc).total_seconds() < 20 \
            and float(truoc[0].diem) == round(data.diem, 1) and truoc[0].so_cau == data.so_cau:
        return _kq_dict(truoc[0], b.tieu_de)                # gửi lặp trong vài giây → không tạo lượt thứ hai
    nguong = data.diem_dat if data.diem_dat is not None else _f(b.diem_dat)
    dat = data.dat if data.dat is not None else ((data.diem >= nguong) if nguong is not None else None)
    nv = _nv_cua(db, nd)

    def cat(v, n):
        return (v.strip()[:n] or None) if v else None

    r = TrnKetQua(bai_id=bai_id, nguoi_dung_id=nd.id, nhan_vien_id=(nv.id if nv is not None else None),
                  ten_nguoi_thi=_ten_nguoi(db, nd)[:150], ten_khai=cat(data.ten_khai, 150), ma_nv_khai=cat(data.ma_nv_khai, 40),
                  don_vi_khai=cat(data.don_vi_khai, 120), chuc_danh_khai=cat(data.chuc_danh_khai, 120),
                  diem=round(data.diem, 1), so_cau=data.so_cau, so_dung=data.so_dung, diem_dat=nguong, dat=dat,
                  lan_thi=len(truoc) + 1, phien_ban=b.phien_ban or 1, thoi_gian_giay=data.thoi_gian_giay,
                  chi_tiet=data.chi_tiet, nop_luc=now)
    db.add(r)
    db.flush()
    ghi_audit(db, nd.id, "TAO", "trn_ket_qua", r.id, moi={"bai": b.tieu_de, "diem": float(r.diem), "dat": dat, "lan": r.lan_thi})
    db.commit()
    return _kq_dict(r, b.tieu_de)


@router.get("/ket-qua")
def ds_ket_qua(bai_id: int | None = None, nhan_vien_id: int | None = None, nguoi_dung_id: int | None = None,
               cua_toi: bool = False, limit: int = 500, db: Session = Depends(get_db),
               nd: NguoiDung = Depends(yeu_cau(MODULE, "XEM"))):
    """Lượt thi. Người quản lý xem được của mọi người (lọc theo bài / nhân viên); người khác chỉ thấy của mình."""
    ql = _la_quan_ly(db, nd)
    q = db.query(TrnKetQua, TrnBai.tieu_de).join(TrnBai, TrnBai.id == TrnKetQua.bai_id)
    if cua_toi or not ql:
        q = q.filter(TrnKetQua.nguoi_dung_id == nd.id)
    else:
        if nhan_vien_id:
            q = q.filter(TrnKetQua.nhan_vien_id == nhan_vien_id)
        if nguoi_dung_id:
            q = q.filter(TrnKetQua.nguoi_dung_id == nguoi_dung_id)
    if bai_id:
        q = q.filter(TrnKetQua.bai_id == bai_id)
    rows = q.order_by(TrnKetQua.nop_luc.desc(), TrnKetQua.id.desc()).limit(max(1, min(int(limit), 5000))).all()
    return [_kq_dict(r, ten) for r, ten in rows]


@router.get("/ket-qua/{kq_id}")
def chi_tiet_ket_qua(kq_id: int, db: Session = Depends(get_db), nd: NguoiDung = Depends(yeu_cau(MODULE, "XEM"))):
    r = db.get(TrnKetQua, kq_id)
    if r is None or (r.nguoi_dung_id != nd.id and not _la_quan_ly(db, nd)):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Không tìm thấy kết quả")
    b = db.get(TrnBai, r.bai_id)
    return _kq_dict(r, b.tieu_de if b else None, chi_tiet=True)


@router.delete("/ket-qua/{kq_id}")
def xoa_ket_qua(kq_id: int, db: Session = Depends(get_db), nd: NguoiDung = Depends(yeu_cau(MODULE, "QUAN_TRI"))):
    r = db.get(TrnKetQua, kq_id)
    if r is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Không tìm thấy kết quả")
    ghi_audit(db, nd.id, "XOA", "trn_ket_qua", r.id,
              cu={"bai_id": r.bai_id, "nguoi": r.ten_nguoi_thi, "diem": float(r.diem), "lan": r.lan_thi})
    db.delete(r)
    db.commit()
    return {"ok": True}


# ================================================================== BẢNG ĐIỂM TỪNG NHÂN VIÊN
def _bang_diem(db: Session) -> dict:
    """Ma trận nhân viên × bài thi: điểm cao nhất, lần gần nhất, số lần, đạt / chưa. Nhân viên đang làm việc luôn có dòng
    (chưa thi = trống); người dùng không có hồ sơ nhân viên mà đã thi vẫn có dòng riêng."""
    bais = (db.query(TrnBai).filter(TrnBai.loai.in_(LOAI_CO_THI)).order_by(TrnBai.id).all())
    nvs = db.query(NhanVien).filter(NhanVien.trang_thai == "DANG_LAM").order_by(NhanVien.ho_ten).all()
    nv_theo_nd = {n.nguoi_dung_id: n for n in nvs if n.nguoi_dung_id}
    hang: dict = {}

    def dong(key, ten, ma=None, chuc_danh=None, nv_id=None, nd_id=None):
        return hang.setdefault(key, {"key": key, "ten": ten, "ma": ma, "chuc_danh": chuc_danh,
                                     "nhan_vien_id": nv_id, "nguoi_dung_id": nd_id, "o": {}})
    for n in nvs:
        dong(f"nv:{n.id}", n.ho_ten, n.ma, n.chuc_danh, n.id, n.nguoi_dung_id)
    ids = [b.id for b in bais]
    if ids:
        for r in (db.query(TrnKetQua).filter(TrnKetQua.bai_id.in_(ids)).order_by(TrnKetQua.nop_luc, TrnKetQua.id).all()):
            nv = (db.get(NhanVien, r.nhan_vien_id) if r.nhan_vien_id else None) or nv_theo_nd.get(r.nguoi_dung_id)
            if nv is not None:
                h = dong(f"nv:{nv.id}", nv.ho_ten, nv.ma, nv.chuc_danh, nv.id, nv.nguoi_dung_id)
            elif r.nguoi_dung_id:
                h = dong(f"nd:{r.nguoi_dung_id}", r.ten_nguoi_thi or f"ND#{r.nguoi_dung_id}", None, None, None, r.nguoi_dung_id)
            else:
                h = dong(f"ten:{r.ten_nguoi_thi}", r.ten_nguoi_thi or "—")
            o = h["o"].setdefault(str(r.bai_id), {"tot_nhat": None, "gan_nhat": None, "so_lan": 0, "dat": False, "ngay": None})
            d = float(r.diem)
            o["so_lan"] += 1
            o["tot_nhat"] = d if o["tot_nhat"] is None else max(o["tot_nhat"], d)
            o["gan_nhat"] = d
            o["ngay"] = str(r.nop_luc)[:10] if r.nop_luc else None
            if r.dat:
                o["dat"] = True
    for h in hang.values():
        h["so_bai_thi"] = len(h["o"])
        h["so_bai_dat"] = sum(1 for x in h["o"].values() if x["dat"])
    return {"bai": [{"id": b.id, "tieu_de": b.tieu_de, "phien_ban": b.phien_ban or 1, "hien": bool(b.hien),
                     "diem_dat": _f(b.diem_dat)} for b in bais],
            "nhan_vien": sorted(hang.values(), key=lambda h: (h["ten"] or "").lower())}


@router.get("/bang-diem")
def bang_diem(db: Session = Depends(get_db), nd: NguoiDung = Depends(yeu_cau(MODULE, "THAO_TAC"))):
    return _bang_diem(db)


def _excel_bytes(bd: dict, luot: list) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    fill = PatternFill("solid", fgColor="E3EDF6")
    xanh = PatternFill("solid", fgColor="DCFCE7")
    do = PatternFill("solid", fgColor="FEE2E2")
    thin = Side(style="thin", color="C9D3DD")
    bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
    luc = gio_hien_tai().strftime("%d/%m/%Y %H:%M")

    def dau(ws, tieu_de, cot, rong):
        ws["A1"] = f"{tieu_de} — xuất lúc {luc}"
        ws["A1"].font = Font(bold=True, size=13)
        for i, (c, w) in enumerate(zip(cot, rong), 1):
            cell = ws.cell(row=3, column=i, value=c)
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.border = bdr
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[3].height = 42
        ws.freeze_panes = "E4"

    ws = wb.active
    ws.title = "Bảng điểm"
    bai = bd["bai"]
    dau(ws, "TRAINING — BẢNG ĐIỂM NHÂN VIÊN (điểm cao nhất, %)",
        ["STT", "Họ tên", "Mã NV", "Chức danh"] + [f"{b['tieu_de']} (Rev {b['phien_ban']})" for b in bai] + ["Số bài đã thi", "Số bài đạt"],
        [6, 28, 12, 24] + [22] * len(bai) + [12, 12])
    for i, h in enumerate(bd["nhan_vien"], 1):
        vals = [i, h["ten"], h["ma"], h["chuc_danh"]]
        for b in bai:
            o = h["o"].get(str(b["id"]))
            vals.append(None if not o else o["tot_nhat"])
        vals += [h["so_bai_thi"], h["so_bai_dat"]]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row=3 + i, column=j, value=v)
            cell.border = bdr
            if 5 <= j < 5 + len(bai) and v is not None:
                o = h["o"].get(str(bai[j - 5]["id"]))
                cell.number_format = "0.0"
                cell.fill = xanh if o and o["dat"] else do
    ws.auto_filter.ref = f"A3:{get_column_letter(4 + len(bai) + 2)}{3 + len(bd['nhan_vien'])}"
    ws2 = wb.create_sheet("Lượt thi")
    dau(ws2, "TRAINING — TẤT CẢ LƯỢT THI",
        ["STT", "Ngày giờ nộp", "Người thi (tài khoản)", "Tên khai trong bài", "Mã NV khai", "Bài", "Rev", "Lần thi", "Điểm (%)",
         "Số câu", "Số đúng", "Điểm đạt (%)", "Kết quả", "Thời gian (giây)"],
        [6, 17, 28, 26, 12, 40, 6, 8, 10, 9, 9, 11, 14, 12])
    for i, r in enumerate(luot, 1):
        vals = [i, r["nop_luc"], r["ten"], r["ten_khai"], r["ma_nv_khai"], r["bai"], r["phien_ban"], r["lan_thi"], r["diem"],
                r["so_cau"], r["so_dung"], r["diem_dat"], ("Đạt" if r["dat"] else ("Chưa đạt" if r["dat"] is False else "")),
                r["thoi_gian_giay"]]
        for j, v in enumerate(vals, 1):
            cell = ws2.cell(row=3 + i, column=j, value=v)
            cell.border = bdr
            if j == 13 and v:
                cell.fill = xanh if v == "Đạt" else do
    ws2.freeze_panes = "A4"
    if luot:
        ws2.auto_filter.ref = f"A3:N{3 + len(luot)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/xuat-excel")
def xuat_excel(db: Session = Depends(get_db), nd: NguoiDung = Depends(yeu_cau(MODULE, "THAO_TAC"))):
    """📥 Excel: sheet «Bảng điểm» (nhân viên × bài) và «Lượt thi» (mọi lượt)."""
    bd = _bang_diem(db)
    rows = (db.query(TrnKetQua, TrnBai.tieu_de).join(TrnBai, TrnBai.id == TrnKetQua.bai_id)
            .order_by(TrnKetQua.nop_luc.desc(), TrnKetQua.id.desc()).limit(20000).all())
    data = _excel_bytes(bd, [_kq_dict(r, ten) for r, ten in rows])
    ten = f"Training-bang-diem-{gio_hien_tai():%Y%m%d}.xlsx"
    return Response(content=data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{ten}"', "Cache-Control": "no-store"})
