"""📋 MẪU BOQ dự toán (theo file «Tab 12 - BOQ-MTO.xlsx» của anh Hiếu, 10/10/2026):
3 dòng đầu (công ty · tên dự toán · tiêu đề bảng), dòng tiêu đề cột
  Mã | Hạng mục | Quy cách | ĐVT | KL | Nhà sản xuất | Nguồn | Đơn giá (VND) | Thành tiền (VND) | Căn cứ tính
dòng NHÓM (Mã một chữ cái, Hạng mục = tên nhóm, không KL), dòng hạng mục (Mã A01…), cuối bảng CỘNG TRƯỚC THUẾ · VAT · TỔNG CỘNG.
 · tao_mau(): sinh file mẫu (trống hoặc điền sẵn các dòng của một dự toán → xuất ra sửa rồi nạp lại)
 · doc_boq(): đọc file theo mẫu THẲNG bằng máy (không AI) → (nhom, items); không đúng mẫu → None để nơi gọi nhờ AI."""
import io
import re
import unicodedata
from datetime import date

COT = ["Mã", "Hạng mục", "Quy cách", "ĐVT", "KL", "Nhà sản xuất", "Nguồn", "Đơn giá (VND)", "Thành tiền (VND)", "Căn cứ tính"]
RONG = [6, 36, 30, 7, 6, 26, 10, 15, 17, 32]
_MAU_TIEU_DE = "E3EDF6"
_MAU_NHOM = "F0F5FA"
TU_KHOA = {
    "ma": ("ma",), "ten": ("hang muc", "ten hang", "ten san pham", "ten", "mo ta hang", "item", "description"),
    "quy_cach": ("quy cach", "spec", "thong so", "mo ta"), "don_vi": ("dvt", "don vi", "unit"),
    "so_luong": ("kl", "sl", "so luong", "khoi luong", "qty", "quantity"),
    "nha_san_xuat": ("nha san xuat", "nsx", "hang san xuat", "brand", "manufacturer"),
    "nguon": ("nguon",), "don_gia": ("don gia", "unit price", "gia"), "thanh_tien": ("thanh tien", "amount", "total"),
    "can_cu": ("can cu", "ghi chu", "note", "remark"), "ncc_ten": ("nha cung cap", "ncc", "supplier", "vendor"),
    "ma_sp": ("ma sp", "ma san pham", "part no", "part number", "model"),
}
_THU_TU = ("ma_sp", "ncc_ten", "nha_san_xuat", "thanh_tien", "don_gia", "so_luong", "don_vi", "quy_cach", "can_cu", "nguon", "ten", "ma")


def _kd(s) -> str:
    s = unicodedata.normalize("NFD", str(s or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s.replace("đ", "d").replace("Đ", "D").lower()).strip()


def _so(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.\-]", "", str(v))
    if not s:
        return 0.0
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        phan = s.split(",")
        s = s.replace(",", "") if (len(phan) > 2 or len(phan[-1]) == 3) else s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[-1]) == 3):
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _tim_tieu_de(o: list) -> dict | None:
    m = {}
    for idx, cell in enumerate(o):
        t = _kd(cell)
        if not t or len(t) > 30:
            continue
        for f in _THU_TU:
            if f in m:
                continue
            if any(t == k or t.startswith(k + " ") or t.startswith(k + "(") or t.startswith(k + "/") for k in TU_KHOA[f]):
                m[f] = idx
                break
    if "ten" not in m or "so_luong" not in m:
        return None
    return m


def doc_boq(data: bytes, filename: str):
    """File Excel / CSV theo mẫu → (nhom: [{ma, ten}], items: [{ma, ten, quy_cach, don_vi, so_luong, nha_san_xuat, nguon,
    don_gia, can_cu, ncc_ten, ma_sp, nhom_ma, nhom_ten}]) hoặc None nếu không nhận diện được tiêu đề cột."""
    fn = (filename or "").lower()
    cac = []
    try:
        if fn.endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                tot = None
                for ws in wb.worksheets:
                    rows = [list(r) for _, r in zip(range(5000), ws.iter_rows(values_only=True))]
                    kq = _doc_bang(rows)
                    if kq and (tot is None or len(kq[1]) > len(tot[1])):
                        tot = kq
                return tot
            finally:
                wb.close()
        if fn.endswith((".csv", ".txt")):
            import csv
            txt = None
            for enc in ("utf-8-sig", "utf-16", "cp1258", "cp1252"):
                try:
                    txt = data.decode(enc)
                    break
                except Exception:
                    continue
            if txt is None:
                return None
            try:
                dialect = csv.Sniffer().sniff(txt[:4000], delimiters=",;\t|")
            except Exception:
                dialect = csv.excel
            cac = [list(r) for r in csv.reader(io.StringIO(txt), dialect)]
            return _doc_bang(cac)
    except Exception:
        return None
    return None


def _doc_bang(cac: list):
    td, i_td = None, -1
    for i, o in enumerate(cac[:40]):
        m = _tim_tieu_de(o)
        if m:
            td, i_td = m, i
            break
    if td is None:
        return None

    def lay(o, f):
        i2 = td.get(f)
        v = o[i2] if (i2 is not None and i2 < len(o)) else None
        return "" if v is None else str(v).strip()

    nhom, items = [], []
    nhom_hien = {"ma": "", "ten": ""}
    for o in cac[i_td + 1:]:
        ma = lay(o, "ma")
        ten = lay(o, "ten")
        if not ten and not ma:
            continue
        tk = _kd(ten)
        if tk.startswith(("cong truoc thue", "tong cong", "vat", "thue gtgt", "cong ", "tong ")) and not lay(o, "so_luong"):
            continue
        sl = _so(lay(o, "so_luong"))
        la_nhom = (bool(re.fullmatch(r"[A-Za-z]{1,2}", ma)) or (not ma and not lay(o, "don_vi"))) and sl <= 0 and not _so(lay(o, "don_gia"))
        if la_nhom:
            if ten:
                nhom_hien = {"ma": ma.upper(), "ten": ten}
                nhom.append(dict(nhom_hien))
            continue
        if not ten:
            continue
        items.append({"ma": ma, "ten": ten[:250], "quy_cach": lay(o, "quy_cach")[:2000], "don_vi": lay(o, "don_vi")[:40],
                      "so_luong": sl, "nha_san_xuat": lay(o, "nha_san_xuat")[:150], "nguon": lay(o, "nguon")[:60],
                      "don_gia": _so(lay(o, "don_gia")), "can_cu": lay(o, "can_cu")[:300],
                      "ncc_ten": lay(o, "ncc_ten")[:200], "ma_sp": lay(o, "ma_sp")[:60],
                      "nhom_ma": nhom_hien["ma"], "nhom_ten": nhom_hien["ten"]})
    if not items:
        return None
    return nhom, items


def ghi_chu_boq(it: dict) -> str:
    """Ghi chú của dòng dự toán giữ mã BOQ · nhóm · nguồn · căn cứ — để xuất lại đúng mẫu."""
    phan = []
    if it.get("ma"):
        phan.append(f"BOQ {it['ma']}")
    if it.get("nhom_ma") or it.get("nhom_ten"):
        phan.append(f"nhóm {it.get('nhom_ma') or ''} {it.get('nhom_ten') or ''}".strip())
    if it.get("nguon"):
        phan.append(f"Nguồn: {it['nguon']}")
    if it.get("can_cu"):
        phan.append(f"Căn cứ: {it['can_cu']}")
    return " · ".join(phan)[:300]


def _tach_ghi_chu(gc: str) -> dict:
    gc = gc or ""
    m_ma = re.search(r"BOQ\s+([A-Za-z]{1,2}\d{1,3})", gc)
    m_nh = re.search(r"nhóm\s+([A-Za-z]{1,2})\s*([^·]*)", gc)
    m_ng = re.search(r"Nguồn:\s*([^·]*)", gc)
    m_cc = re.search(r"Căn cứ:\s*(.*)$", gc)
    return {"ma": m_ma.group(1).upper() if m_ma else "", "nhom_ma": (m_nh.group(1).upper() if m_nh else ""),
            "nhom_ten": (m_nh.group(2).strip() if m_nh else ""), "nguon": (m_ng.group(1).strip() if m_ng else ""),
            "can_cu": (m_cc.group(1).strip() if m_cc else (gc.strip() if gc and "BOQ" not in gc else ""))}


def tao_mau(cong_ty: str, tieu_de: str, ma: str | None = None, items: list | None = None) -> bytes:
    """Sinh file mẫu. items (tùy chọn) = các dòng dự toán hiện có [{ten, quy_cach, don_vi, so_luong, nha_san_xuat, don_gia,
    ghi_chu}] → điền sẵn theo nhóm lưu trong ghi chú (dòng không có nhóm → nhóm «Z HẠNG MỤC KHÁC»)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = (re.sub(r"[^A-Za-z0-9_-]", "-", ma or "BOQ")[:24] + "_BOQ")[:31]
    vien = Side(style="thin", color="B8C4CC")
    khung = Border(left=vien, right=vien, top=vien, bottom=vien)
    ws["A1"] = cong_ty
    ws["A1"].font = Font(bold=True, size=12)
    ws.merge_cells("A1:F1")
    ws["A2"] = tieu_de
    ws["A2"].font = Font(bold=True, size=11)
    ws.merge_cells("A2:F2")
    ws["A3"] = f"BẢNG KHỐI LƯỢNG & DỰ TOÁN — {ma or '<mã dự toán>'} · lập ngày {date.today():%d/%m/%Y}"
    ws.merge_cells("A3:F3")
    for j, c in enumerate(COT, 1):
        cell = ws.cell(row=4, column=j, value=c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=_MAU_TIEU_DE)
        cell.border = khung
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = RONG[j - 1]
    r = 5
    nhoms = []
    if items:
        # gom theo nhóm ghi trong ghi chú; giữ thứ tự xuất hiện
        for it in items:
            t = _tach_ghi_chu(it.get("ghi_chu") or "")
            key = t["nhom_ma"] or "Z"
            g = next((x for x in nhoms if x["ma"] == key), None)
            if g is None:
                g = {"ma": key, "ten": t["nhom_ten"] or ("HẠNG MỤC KHÁC" if key == "Z" else key), "rows": []}
                nhoms.append(g)
            g["rows"].append((it, t))
        nhoms.sort(key=lambda g: (g["ma"] == "Z", g["ma"]))
    else:
        nhoms = [{"ma": "A", "ten": "THIẾT BỊ CÔNG NGHỆ", "rows": [
                    ({"ten": "Ví dụ: Cột lọc áp lực Inox SUS304", "quy_cach": "Ø500×1800 · t=3 mm", "don_vi": "cái", "so_luong": 1,
                      "nha_san_xuat": "SVWS gia công", "don_gia": 0, "ghi_chu": "Nguồn: Đề xuất · Căn cứ: RFQ"}, None),
                    ({"ten": "Ví dụ: Van tự động đa cửa", "quy_cach": "Clack WS1 TC 1\"", "don_vi": "bộ", "so_luong": 2,
                      "nha_san_xuat": "Clack (Mỹ)", "don_gia": 0, "ghi_chu": "Nguồn: Đề xuất"}, None)]},
                 {"ma": "B", "ten": "THIẾT BỊ ĐO & VAN TỰ ĐỘNG", "rows": []},
                 {"ma": "C", "ten": "TỦ ĐIỆN, ĐIỀU KHIỂN & CÁP", "rows": []},
                 {"ma": "D", "ten": "ĐƯỜNG ỐNG, PHỤ KIỆN & VAN TAY", "rows": []},
                 {"ma": "E", "ten": "KẾT CẤU THÉP & GIÁ ĐỠ", "rows": []},
                 {"ma": "F", "ten": "XÂY DỰNG", "rows": []},
                 {"ma": "G", "ten": "CHI PHÍ LẮP ĐẶT", "rows": []}]
    dau_hang = r
    for g in nhoms:
        ws.cell(row=r, column=1, value=g["ma"]).font = Font(bold=True)
        ws.cell(row=r, column=2, value=g["ten"]).font = Font(bold=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        for j in range(1, 11):
            ws.cell(row=r, column=j).fill = PatternFill("solid", fgColor=_MAU_NHOM)
            ws.cell(row=r, column=j).border = khung
        r += 1
        for k, (it, t) in enumerate(g["rows"], 1):
            t = t or _tach_ghi_chu(it.get("ghi_chu") or "")
            ma_dong = t.get("ma") or f"{g['ma']}{k:02d}"
            vals = [ma_dong, it.get("ten") or "", it.get("quy_cach") or "", it.get("don_vi") or "",
                    float(it.get("so_luong") or 0), it.get("nha_san_xuat") or "", t.get("nguon") or "",
                    float(it.get("don_gia") or 0), f"=E{r}*H{r}", t.get("can_cu") or ""]
            for j, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=j, value=v)
                c.border = khung
                c.alignment = Alignment(vertical="top", wrap_text=(j in (2, 3, 6, 10)))
            ws.cell(row=r, column=8).number_format = "#,##0"
            ws.cell(row=r, column=9).number_format = "#,##0"
            r += 1
        if not g["rows"] and not items:
            for j in range(1, 11):
                ws.cell(row=r, column=j).border = khung
            ws.cell(row=r, column=1, value=f"{g['ma']}01")
            ws.cell(row=r, column=9, value=f"=E{r}*H{r}").number_format = "#,##0"
            r += 1
    cuoi_hang = r - 1
    for nhan in ("CỘNG TRƯỚC THUẾ", "VAT 8%", "TỔNG CỘNG"):
        ct = (f"=SUM(I{dau_hang}:I{cuoi_hang})" if nhan.startswith("CỘNG") else
              (f"=I{r - 1}*0.08" if nhan.startswith("VAT") else f"=I{r - 2}+I{r - 1}"))   # công thức theo đúng hàng hiện tại
        ws.cell(row=r, column=2, value=nhan).font = Font(bold=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        c = ws.cell(row=r, column=9, value=ct)
        c.font = Font(bold=True)
        c.number_format = "#,##0"
        for j in range(1, 11):
            ws.cell(row=r, column=j).border = khung
        r += 1
    ws.cell(row=r + 1, column=1, value="Hướng dẫn: giữ nguyên dòng tiêu đề cột (dòng 4). Dòng NHÓM: cột Mã một chữ cái (A, B…), "
                                         "Hạng mục = tên nhóm, để trống KL. Dòng hạng mục: Mã A01, A02…; Thành tiền tự tính. "
                                         "Nạp lại vào app: Bán hàng › Dự toán › 📎 Nạp thêm từ file (máy đọc thẳng, không qua AI).")
    ws.cell(row=r + 1, column=1).font = Font(italic=True, color="5E7178")
    ws.freeze_panes = "A5"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
