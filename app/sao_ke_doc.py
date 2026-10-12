"""Đọc file SAO KÊ ngân hàng dạng bảng (Excel .xlsx / CSV / TXT) THẲNG bằng máy — không qua AI:
đủ mọi dòng, đúng từng đồng, không giới hạn 300 dòng, không tốn phí. Tự tìm dòng tiêu đề (cột Ngày · Diễn giải ·
Ghi nợ / Tiền ra · Ghi có / Tiền vào · Số dư, hoặc một cột Số tiền + cột Loại / dấu). Không nhận diện được thì trả None
để nơi gọi nhờ AI đọc như cũ (PDF / ảnh luôn đi qua AI)."""
import csv
import io
import re
import unicodedata
from datetime import date, datetime

TU_KHOA = {
    "dien_giai": ("dien giai", "noi dung giao dich", "noi dung", "mo ta", "description", "narrative", "remark", "chi tiet giao dich", "content"),
    "so_du": ("so du cuoi", "so du sau gd", "so du", "balance", "running balance"),
    "tien_ra": ("ghi no", "phat sinh no", "tien ra", "so tien ghi no", "debit", "withdrawal", "rut tien", "chi"),
    "tien_vao": ("ghi co", "phat sinh co", "tien vao", "so tien ghi co", "credit", "deposit", "nop tien", "thu"),
    "so_tien": ("so tien giao dich", "so tien", "amount", "gia tri"),
    "loai": ("loai giao dich", "loai gd", "dr/cr", "no/co", "loai", "type"),
    "doi_tuong": ("ten doi tuong", "doi tuong", "nguoi nhan", "ben thu huong", "beneficiary", "counterparty", "doi tac"),
    "ly_do": ("ly do thu/chi", "ly do thu chi", "ly do"),
    "so_ct": ("so chung tu", "so ct", "so tham chieu", "reference", "ma giao dich"),
    "ngay": ("ngay giao dich", "ngay hieu luc", "ngay hach toan", "ngay gd", "transaction date", "value date", "txn date", "ngay", "date"),
}
_THU_TU = ("dien_giai", "so_du", "tien_ra", "tien_vao", "so_tien", "loai", "doi_tuong", "ly_do", "so_ct", "ngay")


def _kd(s) -> str:
    s = unicodedata.normalize("NFD", str(s or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s.replace("đ", "d").replace("Đ", "D").lower()).strip()


def _tim_tieu_de(o: list) -> dict | None:
    """Ô của một dòng → {field: chỉ số cột} nếu dòng này là tiêu đề bảng giao dịch."""
    map_ = {}
    for idx, cell in enumerate(o):
        t = _kd(cell)
        if not t or len(t) > 40:
            continue
        for f in _THU_TU:
            if f in map_:
                continue
            if any(t == k or t.startswith(k + " ") or t.startswith(k + "(") or (len(k) >= 6 and k in t) for k in TU_KHOA[f]):
                map_[f] = idx
                break
    if "ngay" not in map_:
        return None
    if not (("tien_ra" in map_ or "tien_vao" in map_) or "so_tien" in map_):
        return None
    return map_


def _ngay(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    if not s:
        return None
    m = re.search(r"(?<!\d)(\d{4})[/\-.](\d{1,2})[/\-.](\d{1,2})(?!\d)", s)       # 2026-08-01 (xét trước, kẻo 26-08-01 bị hiểu nhầm)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    m = re.search(r"(?<!\d)(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4}|\d{2})(?!\d)", s)  # 01/08/2026 · 01-08-26
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y < 100:
            y += 2000
        try:
            return date(y, mo, d)
        except ValueError:
            return None
    return None


def _so(v) -> float:
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s:
        return 0.0
    am = s.startswith("(") and s.endswith(")") or s.startswith("-")
    s = re.sub(r"[^\d,.\-]", "", s)
    if not s:
        return 0.0
    # 1.234.567 / 1,234,567 / 1.234.567,00 / 1,234,567.00 / 1234567.5
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):          # 1.234.567,00 → phần thập phân sau dấu phẩy
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        phan = s.split(",")
        s = s.replace(",", "") if (len(phan) > 2 or len(phan[-1]) == 3) else s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[-1]) == 3):
        s = s.replace(".", "")
    try:
        x = float(s)
    except ValueError:
        return 0.0
    return -abs(x) if am else x


_RE_RA = re.compile(r"\b(uy nhiem chi|sec|phieu chi|bao no|debit|ghi no|chuyen khoan di|rut tien)\b")
_RE_VAO = re.compile(r"\b(thu tien|bao co|phieu thu|credit|ghi co|nop tien|deposit|thu khac)\b")


def _chieu_theo_loai(ly_do: str, loai: str, dien_giai: str):
    """Bản xuất kế toán (CUKCUK…) có MỘT cột Số tiền dương — chiều nằm ở «Lý do thu/chi» (Chi khác · Thu khác · Chuyển tiền
    nội bộ) và «Loại chứng từ» (Ủy nhiệm chi · Séc tiền mặt · Thu tiền gửi…). Trả "RA" | "VAO" | None (không quyết được).
    Chuyển tiền nội bộ: «mua ngoại tệ» là tiền ra, còn lại tính tiền vào (chiều chính xác chỉ nhìn được qua dòng tổng cuối file)."""
    ld = (ly_do or "").strip()
    if ld.startswith("chi"):
        return "RA"
    if ld.startswith("thu"):
        return "VAO"
    t = (str(loai or "") + " " + ld).strip()
    if _RE_RA.search(t):
        return "RA"
    if _RE_VAO.search(t):
        return "VAO"
    if "noi bo" in t:
        return "RA" if "ngoai te" in (dien_giai or "") else "VAO"
    return None


def _tong_cuoi_file(cac_dong: list[list]) -> dict:
    """Dòng tổng cuối file («Số dòng = 334 · Tổng thu: … · Tổng chi: …») → để đối chiếu với số đọc được."""
    kq = {}
    for o in cac_dong:
        for c in o:
            if not isinstance(c, str):
                continue
            t = _kd(c)
            m = re.search(r"tong thu\s*:?\s*([\d.,]+)", t)
            if m:
                kq["tong_thu_file"] = _so(m.group(1))
            m = re.search(r"tong chi\s*:?\s*([\d.,]+)", t)
            if m:
                kq["tong_chi_file"] = _so(m.group(1))
            m = re.search(r"so dong\s*=?\s*(\d+)", t)
            if m:
                kq["so_dong_file"] = int(m.group(1))
    return kq


def _bang_tu_o(cac_dong: list[list]) -> tuple[list[dict], dict] | None:
    """Danh sách dòng (mỗi dòng là list ô) → (giao dịch, meta) hoặc None khi không thấy tiêu đề."""
    tieu_de, i_td = None, -1
    for i, o in enumerate(cac_dong[:60]):
        m = _tim_tieu_de(o)
        if m:
            tieu_de, i_td = m, i
            break
    if tieu_de is None:
        return None
    ra = []
    dau = tieu_de.get("loai")
    noi_bo_idx = []          # chỉ số các dòng chuyển tiền nội bộ (chiều suy từ diễn giải) — chỉnh theo tổng cuối file nếu lệch
    for o in cac_dong[i_td + 1:]:
        def lay(f):
            i2 = tieu_de.get(f)
            return o[i2] if (i2 is not None and i2 < len(o)) else None
        ng = _ngay(lay("ngay"))
        dg = str(lay("dien_giai") or "").strip()
        vao = _so(lay("tien_vao")) if "tien_vao" in tieu_de else 0.0
        ra_ = _so(lay("tien_ra")) if "tien_ra" in tieu_de else 0.0
        la_noi_bo = False
        if "so_tien" in tieu_de and not vao and not ra_:
            st = _so(lay("so_tien"))
            lo = _kd(lay("loai")) if dau is not None else ""
            ld = _kd(lay("ly_do")) if "ly_do" in tieu_de else ""
            ch = _chieu_theo_loai(ld, lo, _kd(dg)) if (ld or lo) else None
            la_noi_bo = "noi bo" in (lo + " " + ld)
            if ch == "RA":
                ra_ = abs(st)
            elif ch == "VAO":
                vao = abs(st)
            elif st < 0 or lo.startswith(("no", "d", "-", "debit", "ra", "chi", "withdraw")):
                ra_ = abs(st)
            elif st > 0:
                vao = st
        vao, ra_ = abs(vao), abs(ra_)
        if ng is None:
            # dòng nối tiếp diễn giải (ngân hàng xuống dòng nội dung) → ghép vào dòng trước
            if ra and dg and not vao and not ra_:
                ra[-1]["dien_giai"] = (ra[-1]["dien_giai"] + " " + dg).strip()[:400]
            continue
        if not vao and not ra_:
            continue                       # số dư đầu kỳ / dòng tổng
        sd = lay("so_du")
        dt_ten = str(lay("doi_tuong") or "").strip() if "doi_tuong" in tieu_de else ""
        if dt_ten and _kd(dt_ten) not in _kd(dg):
            dg = (dg + " - " + dt_ten).strip(" -")      # cùng dạng «TT HD 17 - TÊN ĐỐI TÁC» để nhận diện đối tác từ diễn giải
        so_ct = str(lay("so_ct") or "").strip()[:60] if "so_ct" in tieu_de else ""
        if la_noi_bo:
            noi_bo_idx.append(len(ra))
        ra.append({"ngay": ng, "dien_giai": dg[:400], "tien_vao": vao, "tien_ra": ra_, "so_ct": so_ct or None,
                   "so_du": _so(sd) if (sd not in (None, "")) else None})
    meta = {"cach": "MAY", "dong_tieu_de": i_td + 1, "cot": tieu_de, "noi_bo": len(noi_bo_idx)}
    meta.update(_tong_cuoi_file(cac_dong[i_td + 1:]))
    # 🎯 Đối chiếu dòng tổng cuối file: lệch đúng bằng MỘT dòng chuyển nội bộ → dòng đó mang chiều ngược lại
    tt, tc = meta.get("tong_thu_file"), meta.get("tong_chi_file")
    if tt is not None and tc is not None and noi_bo_idx:
        sv, sr = sum(r["tien_vao"] for r in ra), sum(r["tien_ra"] for r in ra)
        dv = sv - tt
        if abs(dv) > 1 and abs((sr + dv) - tc) <= 1:
            for i in noi_bo_idx:
                r = ra[i]
                if dv > 0 and r["tien_vao"] and abs(r["tien_vao"] - dv) <= 0.5:
                    r["tien_ra"], r["tien_vao"] = r["tien_vao"], 0.0
                    meta["da_chinh_chieu"] = 1
                    break
                if dv < 0 and r["tien_ra"] and abs(r["tien_ra"] + dv) <= 0.5:
                    r["tien_vao"], r["tien_ra"] = r["tien_ra"], 0.0
                    meta["da_chinh_chieu"] = 1
                    break
    return ra, meta


def doc_xlsx(data: bytes):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        tot = None
        for ws in wb.worksheets:
            cac = []
            for row in ws.iter_rows(values_only=True):
                cac.append(list(row))
                if len(cac) >= 20000:
                    break
            kq = _bang_tu_o(cac)
            if kq and (tot is None or len(kq[0]) > len(tot[0])):
                kq[1]["sheet"] = ws.title
                tot = kq
        return tot
    finally:
        wb.close()


def doc_csv(data: bytes):
    txt = None
    for enc in ("utf-8-sig", "utf-16", "cp1258", "cp1252"):
        try:
            txt = data.decode(enc)
            break
        except Exception:
            continue
    if txt is None:
        return None
    mau = txt[:5000]
    try:
        dialect = csv.Sniffer().sniff(mau, delimiters=",;\t|")
    except Exception:
        dialect = csv.excel
    cac = [list(r) for r in csv.reader(io.StringIO(txt), dialect)]
    return _bang_tu_o(cac)


def doc_bang(data: bytes, content_type: str, filename: str):
    """Excel / CSV / TXT → (giao dịch, meta) hoặc None (không phải bảng, hoặc không nhận diện được cột)."""
    fn = (filename or "").lower()
    ct = (content_type or "").lower()
    try:
        if fn.endswith((".xlsx", ".xlsm")) or ct == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            return doc_xlsx(data)
        if fn.endswith((".csv", ".txt")) or ct in ("text/csv", "text/plain"):
            return doc_csv(data)
    except Exception:
        return None
    return None
